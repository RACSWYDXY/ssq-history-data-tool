"""Offline tests: no network, browser or remote history download is used."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

import ssq_downloader as ssq


def row_html(issue, draw, machine):
    cells = [
        issue, '2026-10-25', '周日', machine, '100', '24', '3:3', '3:3',
        draw, '99', '27', '4:2', '3:3',
    ]
    return '<tr>' + ''.join(f'<td>{value}</td>' for value in cells) + '</tr>'


def page_html(*rows):
    return (
        '<html><table><tr><th>期号</th><th>开奖日期</th>'
        '<th>开机号</th><th>奖号</th></tr>'
        + ''.join(rows) + '</table></html>'
    )


def fixture_records():
    return ssq.parse_page(page_html(
        row_html('2026122', '30 02 19 01 17 04 + 09', '18 08 25 16 03 23 + 05'),
        row_html('2026121', '13 18 14 20 31 16 + 03', '31 01 17 09 24 33 + 11'),
    ))


def test_parse_unsorted_reds_keep_original_fields_and_order_issues():
    rows = fixture_records()
    assert [r['期号'] for r in rows] == ['2026121', '2026122']
    latest = rows[-1]
    assert latest['奖号红球'] == [1, 2, 4, 17, 19, 30]
    assert latest['开机红球'] == [3, 8, 16, 18, 23, 25]
    assert latest['奖号号码'] == '30 02 19 01 17 04 + 09'
    assert latest['奖号蓝球'] == 9


def test_derived_values_and_overlap():
    rows = fixture_records()
    row = ssq.derived_draw_row(rows[1], rows[0]['奖号红球'])
    assert row[3:9] == [1, 2, 4, 17, 19, 30]
    assert row[9] == 9
    assert row[10] == 73  # sum
    assert row[11] == 29  # span
    assert row[12:14] == [3, 3]
    assert row[16] == 1  # 1,2 as one consecutive group
    assert row[17] == 0
    assert row[18] == '3:2:1'


def test_excel_export_sheets_and_frequency(tmp_path):
    rows = fixture_records()
    destination = tmp_path / 'sample.xlsx'
    ssq.write_excel(rows, destination)
    wb = load_workbook(destination, read_only=True)
    assert wb.sheetnames == ['开奖记录', '开机号', '原始数据', '号码频率', '说明']
    assert [wb['开奖记录'].cell(3, col).value for col in range(4, 10)] == [1, 2, 4, 17, 19, 30]
    assert wb['原始数据'].cell(3, 9).value == '30 02 19 01 17 04 + 09'
    assert wb['号码频率'].cell(2, 2).value == 1  # red 01 occurs once
    assert wb['号码频率'].cell(2, 3).value == 1
    assert wb['号码频率'].cell(2, 9).value == 0  # blue 01 absent
    wb.close()


def test_validation_rejects_duplicate_red_balls(monkeypatch):
    monkeypatch.setattr(ssq, 'MIN_RECORDS', 1)
    rows = fixture_records()
    rows[-1]['奖号红球'] = [1, 1, 4, 17, 19, 30]
    with pytest.raises(RuntimeError, match='非法'):
        ssq.check_records(rows)


def test_validation_rejects_short_history():
    with pytest.raises(RuntimeError, match='少于完整历史'):
        ssq.check_records(fixture_records())
