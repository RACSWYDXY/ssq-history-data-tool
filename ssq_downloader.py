# -*- coding: utf-8 -*-
"""
双色球历史数据下载器
数据源：https://www.17500.cn/chart/ssq-kjh.html

功能：
1. 自动打开网页并尝试切换到“4000期”（足够覆盖双色球历史期数）
2. 提取期号、日期、开机号、奖号及网页统计字段
3. 拆分 6 个红球 + 1 个蓝球
4. 导出 Excel / CSV / TXT
5. 生成基础的描述性频率统计（不做预测）

首次运行依赖：selenium, beautifulsoup4, openpyxl
建议使用同目录下的“run.bat”。
"""

from __future__ import annotations

import csv
import re
import sys
import time
import traceback
import logging
import requests
from urllib.parse import urljoin
from collections import Counter
from datetime import datetime
from pathlib import Path

URL = "https://www.17500.cn/chart/ssq-kjh.html"
MIN_RECORDS = 3000  # 2003年以来的历史记录；不允许只有几十或几百期时冒充完整数据
OUTPUT_DIR = Path(__file__).resolve().parent / "双色球数据"


def check_dependencies():
    missing = []
    try:
        import requests  # noqa: F401
    except ImportError:
        missing.append("requests")
    try:
        import selenium  # noqa: F401
    except ImportError:
        missing.append("selenium")
    try:
        import bs4  # noqa: F401
    except ImportError:
        missing.append("beautifulsoup4")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        missing.append("openpyxl")

    if missing:
        print("缺少 Python 依赖：" + ", ".join(missing))
        print("请双击同目录下的 run.bat，它会自动安装依赖。")
        sys.exit(1)


from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ISSUE_RE = re.compile(r"^20\d{5}$")


def make_driver():
    """尝试 Chrome/Edge 的普通及无头模式；记录每一种失败详情。"""
    from selenium import webdriver

    errors = []
    for browser in ('Edge', 'Chrome'):
        for headless in (True, False):
            try:
                options = getattr(webdriver, browser + 'Options')()
                if headless:
                    options.add_argument('--headless=new')
                options.add_argument('--window-size=1600,1200')
                options.add_argument('--disable-extensions')
                options.add_argument('--no-first-run')
                options.add_argument('--disable-notifications')
                driver = getattr(webdriver, browser)(options=options)
                driver.set_page_load_timeout(55)
                print(f'成功启动 {browser} (无头模式={headless})')
                return driver
            except Exception as exc:
                detail = ''.join(traceback.format_exception_only(type(exc), exc)).strip()
                errors.append(f'{browser}, headless={headless}: {detail}')
                print(f'{browser} 启动失败 (无头={headless})：{detail[:230]}')
    raise RuntimeError('所有浏览器启动方案均失败：\n' + '\n'.join(errors))


def direct_download():
    """无浏览器获取：先请求默认页，尝试可识别的 4000 单选按钮所在表单。"""
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126.0 Safari/537.36',
        'Accept-Language': 'zh-CN,zh;q=0.9',
    })
    response = session.get(URL, timeout=(12, 35))
    response.raise_for_status()
    response.encoding = response.apparent_encoding or response.encoding
    html = response.text
    initial = parse_page(html)
    print(f'直接访问：解析到 {len(initial)} 期')
    if len(initial) >= MIN_RECORDS:
        return initial
    soup = BeautifulSoup(html, 'html.parser')
    form_attempts = []
    for inp in soup.select('input[value="4000"]'):
        form = inp.find_parent('form')
        if form is None:
            continue
        payload = {}
        for field in form.find_all(['input', 'select']):
            name = field.get('name')
            if not name:
                continue
            if field.name == 'select':
                opt = field.find('option', selected=True) or field.find('option')
                if opt:
                    payload[name] = opt.get('value', opt.text.strip())
            elif field.get('type', '').lower() in ('radio', 'checkbox'):
                if field is inp or field.has_attr('checked'):
                    payload[name] = field.get('value', 'on')
            elif field.get('type', '').lower() not in ('button', 'reset', 'file'):
                payload[name] = field.get('value', '')
        if inp.get('name'):
            payload[inp['name']] = '4000'
        action = urljoin(response.url, form.get('action') or response.url)
        method = (form.get('method') or 'get').lower()
        form_attempts.append((method, action, payload))
    for method, action, payload in form_attempts[:4]:
        try:
            resp = session.post(action, data=payload, timeout=(12, 40)) if method == 'post' else session.get(action, params=payload, timeout=(12, 40))
            resp.raise_for_status()
            resp.encoding = resp.apparent_encoding or resp.encoding
            rows = parse_page(resp.text)
            print(f'尝试网页表单 {method.upper()}：{len(rows)} 期')
            if len(rows) >= MIN_RECORDS:
                return rows
        except Exception as exc:
            print(f'网页表单请求失败：{type(exc).__name__}: {exc}')
    raise RuntimeError(f'直接抓取仅有 {len(initial)} 期，不能当作完整历史。将尝试浏览器方式。')


def browser_download():
    from selenium.webdriver.support.ui import WebDriverWait

    driver = None
    try:
        driver = make_driver()
        print('正在加载网站（浏览器已启动）……')
        driver.get(URL)
        WebDriverWait(driver, 30).until(lambda d: issue_row_count(d) >= 10)
        click_4000(driver)
        WebDriverWait(driver, 30).until(lambda d: issue_row_count(d) >= MIN_RECORDS)
        html = driver.page_source
        rows = parse_page(html)
        if len(rows) < MIN_RECORDS:
            (OUTPUT_DIR / 'debug_page.html').write_text(html, encoding='utf-8')
            raise RuntimeError(f'只有 {len(rows)} 期。已保存 debug_page.html')
        return rows
    finally:
        if driver is not None:
            driver.quit()

def issue_row_count(driver) -> int:
    from selenium.webdriver.common.by import By

    count = 0
    for tr in driver.find_elements(By.TAG_NAME, "tr"):
        try:
            tds = tr.find_elements(By.TAG_NAME, "td")
            if tds and ISSUE_RE.fullmatch(tds[0].text.strip()):
                count += 1
        except Exception:
            continue
    return count


def click_4000(driver):
    """尽量用多种方式切换到 4000 期，避免依赖网站内部 URL 参数。"""
    from selenium.webdriver.common.by import By

    before = issue_row_count(driver)
    print(f"网页初始可见数据：{before} 期")

    candidates = []
    xpaths = [
        "//input[@value='4000']",
        "//label[normalize-space(.)='4000']",
        "//*[self::span or self::a or self::button][normalize-space(.)='4000']",
    ]
    for xp in xpaths:
        try:
            candidates.extend(driver.find_elements(By.XPATH, xp))
        except Exception:
            pass

    clicked = False
    for el in candidates:
        try:
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
            driver.execute_script("arguments[0].click();", el)
            clicked = True
            break
        except Exception:
            continue

    if clicked:
        # 页面可能局部刷新或整页跳转。
        for _ in range(20):
            time.sleep(0.5)
            now = issue_row_count(driver)
            if now > max(before + 20, 100):
                print(f"已切换到长列表：{now} 期")
                return

        # 某些实现需要再点一次“查看”。
        try:
            buttons = driver.find_elements(
                By.XPATH,
                "//button[contains(normalize-space(.),'查看')] | "
                "//input[@type='submit' and contains(@value,'查看')]",
            )
            if buttons:
                driver.execute_script("arguments[0].click();", buttons[0])
                time.sleep(2)
        except Exception:
            pass

    now = issue_row_count(driver)
    if now <= before + 5:
        # 最后的兜底：尝试点击“全部”。
        try:
            all_els = driver.find_elements(
                By.XPATH,
                "//*[self::a or self::button or self::span][normalize-space(.)='全部']",
            )
            for el in all_els:
                try:
                    driver.execute_script("arguments[0].click();", el)
                    time.sleep(2)
                    if issue_row_count(driver) > now + 20:
                        break
                except Exception:
                    continue
        except Exception:
            pass

    print(f"最终页面可见数据：{issue_row_count(driver)} 期")


def clean_text(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def parse_ball_text(text: str):
    """将 '01 03 19 20 24 25 + 07' 拆成 [1,3,19,20,24,25], 7。"""
    text = clean_text(text)
    if not text:
        return None, None
    parts = re.split(r"\s*\+\s*", text)
    if len(parts) != 2:
        nums = [int(x) for x in re.findall(r"\d{1,2}", text)]
        if len(nums) >= 7:
            return nums[:6], nums[6]
        return None, None

    reds = [int(x) for x in re.findall(r"\d{1,2}", parts[0])]
    blues = [int(x) for x in re.findall(r"\d{1,2}", parts[1])]
    if len(reds) != 6 or len(blues) < 1:
        return None, None
    return reds, blues[0]


def find_target_table(soup: BeautifulSoup):
    best = None
    best_score = -1
    for table in soup.find_all("table"):
        txt = clean_text(table.get_text(" ", strip=True))
        score = sum(k in txt for k in ("期号", "开奖日期", "开机号", "奖号"))
        if score > best_score:
            best, best_score = table, score
    if best_score < 3:
        raise RuntimeError("没有找到包含“期号/开奖日期/开机号/奖号”的目标表格。")
    return best


def parse_page(html: str):
    soup = BeautifulSoup(html, "html.parser")
    table = find_target_table(soup)
    records = []

    for tr in table.find_all("tr"):
        cells = [clean_text(td.get_text(" ", strip=True)) for td in tr.find_all("td")]
        if not cells or not ISSUE_RE.fullmatch(cells[0]):
            continue

        # 正常页面为 13 列：期号、日期、星期 + 开机号5列 + 奖号5列。
        if len(cells) < 13:
            # 少数页面因 rowspan/结构变化，先跳过，避免错列。
            continue

        rec = {
            "期号": cells[0],
            "开奖日期": cells[1],
            "星期": cells[2],
            "开机号码": cells[3],
            "开机和值": cells[4],
            "开机极距": cells[5],
            "开机大小": cells[6],
            "开机奇偶": cells[7],
            "奖号号码": cells[8],
            "奖号和值": cells[9],
            "奖号极距": cells[10],
            "奖号大小": cells[11],
            "奖号奇偶": cells[12],
        }
        rec["开机红球"], rec["开机蓝球"] = parse_ball_text(rec["开机号码"])
        rec["奖号红球"], rec["奖号蓝球"] = parse_ball_text(rec["奖号号码"])

        # 原始字符串留在“原始数据”工作表；导出的分列红球始终按数值升序。
        # 这样网页即使按出球顺序或非标准顺序展示，也不影响红1—红6的含义。
        for key in ("开机红球", "奖号红球"):
            if rec[key] is not None:
                rec[key] = sorted(rec[key])
        records.append(rec)

    # 去重 + 期号升序
    by_issue = {r["期号"]: r for r in records}
    records = sorted(by_issue.values(), key=lambda r: int(r["期号"]))
    return records


def consecutive_groups(nums):
    if not nums:
        return 0
    nums = sorted(nums)
    groups = 0
    in_run = False
    for a, b in zip(nums, nums[1:]):
        if b == a + 1:
            if not in_run:
                groups += 1
                in_run = True
        else:
            in_run = False
    return groups


def derived_draw_row(rec, prev_reds=None):
    reds = rec["奖号红球"]
    blue = rec["奖号蓝球"]
    if not reds or blue is None:
        return None

    odd = sum(n % 2 == 1 for n in reds)
    big = sum(n >= 17 for n in reds)
    overlap_prev = len(set(reds) & set(prev_reds or []))
    z1 = sum(1 <= n <= 11 for n in reds)
    z2 = sum(12 <= n <= 22 for n in reds)
    z3 = sum(23 <= n <= 33 for n in reds)

    return [
        rec["期号"], rec["开奖日期"], rec["星期"],
        *reds, blue,
        sum(reds), max(reds) - min(reds),
        odd, 6 - odd, big, 6 - big,
        consecutive_groups(reds), overlap_prev,
        f"{z1}:{z2}:{z3}",
    ]


def derived_machine_row(rec):
    reds = rec["开机红球"]
    blue = rec["开机蓝球"]
    if not reds or blue is None:
        return None
    draw_reds = rec["奖号红球"] or []
    draw_blue = rec["奖号蓝球"]
    return [
        rec["期号"], rec["开奖日期"], rec["星期"],
        *reds, blue,
        sum(reds), max(reds) - min(reds),
        sum(n % 2 == 1 for n in reds),
        sum(n >= 17 for n in reds),
        len(set(reds) & set(draw_reds)),
        1 if draw_blue is not None and blue == draw_blue else 0,
    ]


def style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            val = "" if cell.value is None else str(cell.value)
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(val) + 2), 24)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = max(width, 9)


def write_excel(records, path: Path):
    wb = Workbook()
    wb.remove(wb.active)

    # 1) 开奖记录
    ws = wb.create_sheet("开奖记录")
    headers = [
        "期号", "开奖日期", "星期",
        "红1", "红2", "红3", "红4", "红5", "红6", "蓝球",
        "红球和值", "红球跨度", "奇数数", "偶数数", "大数数(17-33)", "小数数(1-16)",
        "连号组数", "与上期红球重号数", "三区比(1-11:12-22:23-33)",
    ]
    ws.append(headers)
    prev = None
    for rec in records:
        row = derived_draw_row(rec, prev)
        if row:
            ws.append(row)
            prev = rec["奖号红球"]
    style_sheet(ws)
    for col in range(4, 11):
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = "00"

    # 2) 开机号
    ws2 = wb.create_sheet("开机号")
    headers2 = [
        "期号", "开奖日期", "星期",
        "开机红1", "开机红2", "开机红3", "开机红4", "开机红5", "开机红6", "开机蓝",
        "红球和值", "红球跨度", "奇数数", "大数数(17-33)",
        "与奖号红球重合数", "蓝球是否相同",
    ]
    ws2.append(headers2)
    for rec in records:
        row = derived_machine_row(rec)
        if row:
            ws2.append(row)
    style_sheet(ws2)
    for col in range(4, 11):
        for cell in ws2.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = "00"

    # 3) 原始数据
    ws3 = wb.create_sheet("原始数据")
    raw_headers = [
        "期号", "开奖日期", "星期", "开机号码", "开机和值", "开机极距", "开机大小", "开机奇偶",
        "奖号号码", "奖号和值", "奖号极距", "奖号大小", "奖号奇偶",
    ]
    ws3.append(raw_headers)
    for r in records:
        ws3.append([r[h] for h in raw_headers])
    style_sheet(ws3)
    ws3.column_dimensions["D"].width = 28
    ws3.column_dimensions["I"].width = 28

    # 4) 号码频率：全历史 / 最近100 / 50 / 30期，仅描述性统计
    ws4 = wb.create_sheet("号码频率")
    completed = [r for r in records if r["奖号红球"] and r["奖号蓝球"] is not None]
    windows = [("全历史", completed), ("最近100期", completed[-100:]), ("最近50期", completed[-50:]), ("最近30期", completed[-30:])]

    red_counters = {name: Counter(n for r in rows for n in r["奖号红球"]) for name, rows in windows}
    blue_counters = {name: Counter(r["奖号蓝球"] for r in rows) for name, rows in windows}

    ws4.append(["红球", "全历史次数", "最近100期", "最近50期", "最近30期", "全历史占全部红球位比例"])
    total_red_positions = max(1, len(completed) * 6)
    for n in range(1, 34):
        ws4.append([
            n,
            red_counters["全历史"][n], red_counters["最近100期"][n],
            red_counters["最近50期"][n], red_counters["最近30期"][n],
            red_counters["全历史"][n] / total_red_positions,
        ])
    ws4.cell(1, 8, "蓝球")
    ws4.cell(1, 9, "全历史次数")
    ws4.cell(1, 10, "最近100期")
    ws4.cell(1, 11, "最近50期")
    ws4.cell(1, 12, "最近30期")
    for i, n in enumerate(range(1, 17), start=2):
        ws4.cell(i, 8, n)
        ws4.cell(i, 9, blue_counters["全历史"][n])
        ws4.cell(i, 10, blue_counters["最近100期"][n])
        ws4.cell(i, 11, blue_counters["最近50期"][n])
        ws4.cell(i, 12, blue_counters["最近30期"][n])

    # 自定义频率页格式
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for c in list(ws4[1][:6]) + list(ws4[1][7:12]):
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")
    ws4.freeze_panes = "A2"
    ws4.sheet_view.showGridLines = False
    for c in ws4["A"][1:] + ws4["H"][1:]:
        c.number_format = "00"
    for c in ws4["F"][1:]:
        c.number_format = "0.00%"
    for col in range(1, 13):
        ws4.column_dimensions[get_column_letter(col)].width = 18

    # 5) 说明
    ws5 = wb.create_sheet("说明")
    notes = [
        ["双色球历史数据下载器"],
        ["数据源", URL],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["开奖记录", "仅保留已有正式奖号的期次；红球和蓝球拆成独立列，方便筛选/统计。"],
        ["开机号", "网站展示的模拟开机号，与正式开奖数据分开保存。"],
        ["原始数据", "尽量原样保留网页中的和值、极距、大小、奇偶等字段。"],
        ["号码频率", "仅为历史出现次数的描述性统计，不代表未来开奖概率或预测结果。"],
        ["更新方法", "以后重新运行 run.bat 即可更新同名数据文件。"],
    ]
    for row in notes:
        ws5.append(row)
    ws5["A1"].font = Font(size=16, bold=True)
    ws5.column_dimensions["A"].width = 20
    ws5.column_dimensions["B"].width = 90
    ws5.sheet_view.showGridLines = False

    wb.save(path)


def write_csv(records, path: Path):
    headers = [
        "期号", "开奖日期", "星期",
        "红1", "红2", "红3", "红4", "红5", "红6", "蓝球",
        "红球和值", "红球跨度", "奇数数", "偶数数", "大数数", "小数数", "连号组数", "与上期红球重号数", "三区比",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        prev = None
        for rec in records:
            row = derived_draw_row(rec, prev)
            if row:
                w.writerow(row)
                prev = rec["奖号红球"]


def write_txt(records, draw_path: Path, machine_path: Path):
    with draw_path.open("w", encoding="utf-8") as f:
        for r in records:
            reds, blue = r["奖号红球"], r["奖号蓝球"]
            if reds and blue is not None:
                balls = " ".join(f"{n:02d}" for n in reds)
                f.write(f"{r['期号']} {r['开奖日期']} {balls} + {blue:02d}\n")

    with machine_path.open("w", encoding="utf-8") as f:
        for r in records:
            reds, blue = r["开机红球"], r["开机蓝球"]
            if reds and blue is not None:
                balls = " ".join(f"{n:02d}" for n in reds)
                f.write(f"{r['期号']} {r['开奖日期']} {balls} + {blue:02d}\n")


def check_records(records):
    if len(records) < MIN_RECORDS:
        raise RuntimeError(f'下载到 {len(records)} 期，少于完整历史最低校验值 {MIN_RECORDS}，拒绝覆盖旧数据。')
    completed = [r for r in records if r['奖号红球'] and r['奖号蓝球'] is not None]
    if len(completed) < MIN_RECORDS:
        raise RuntimeError(f'有效开奖数据仅 {len(completed)} 期，拒绝覆盖旧文件。')
    issues = [r['期号'] for r in records]
    if len(set(issues)) != len(issues):
        raise RuntimeError('期号存在重复，拒绝导出')
    for r in completed:
        reds, blue = r['奖号红球'], r['奖号蓝球']
        if len(reds) != 6 or len(set(reds)) != 6 or not all(1 <= x <= 33 for x in reds) or not 1 <= blue <= 16:
            raise RuntimeError(f"{r['期号']} 的号码非法，拒绝导出")
        if reds != sorted(reds):
            raise RuntimeError(f"{r['期号']} 的分列红球不是升序，拒绝导出")
    print(f'通过校验：页面 {len(records)} 期，有效开奖 {len(completed)} 期；最新期号 {issues[-1]}')


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    log_path = OUTPUT_DIR / '运行日志.txt'
    # 日志保留完整 traceback，便于反馈问题
    logging.basicConfig(filename=log_path, level=logging.INFO, encoding='utf-8',
                        format='%(asctime)s %(levelname)s %(message)s', force=True)
    print('=' * 64)
    check_dependencies()
    print('双色球数据下载器 · 数据下载与描述性统计')
    print('来源：', URL)
    print('详细日志：', log_path)
    print('=' * 64)
    errors = []
    try:
        records = None
        print('\n[1/4] 尝试直接访问网页（无需浏览器）……')
        try:
            records = direct_download()
        except Exception as exc:
            errors.append(('直接抓取', traceback.format_exc()))
            print(f'直接获取未完成：{type(exc).__name__}: {exc}')
        if records is None:
            print('\n[2/4] 尝试浏览器自动化作为备用……')
            try:
                records = browser_download()
            except Exception as exc:
                errors.append(('浏览器备用方案', traceback.format_exc()))
                print(f'浏览器方法失败：{type(exc).__name__}: {exc}')
        if records is None:
            raise RuntimeError('两种下载方式都失败。请将「双色球数据/运行日志.txt」发给我以进一步定位。')
        check_records(records)
        print('\n[3/4] 正在导出 Excel / CSV / TXT……')
        # 先写临时文件；成功生成全部格式后才替换旧文件，尽量避免半成品。
        exports = [
            ('双色球历史数据.tmp.xlsx', '双色球历史数据.xlsx'),
            ('双色球开奖记录.tmp.csv', '双色球开奖记录.csv'),
            ('双色球开奖号.tmp.txt', '双色球开奖号.txt'),
            ('双色球开机号.tmp.txt', '双色球开机号.txt'),
        ]
        try:
            write_excel(records, OUTPUT_DIR / exports[0][0])
            write_csv(records, OUTPUT_DIR / exports[1][0])
            write_txt(records, OUTPUT_DIR / exports[2][0], OUTPUT_DIR / exports[3][0])
            for temp_name, final_name in exports:
                (OUTPUT_DIR / temp_name).replace(OUTPUT_DIR / final_name)
        finally:
            for temp_name, _ in exports:
                (OUTPUT_DIR / temp_name).unlink(missing_ok=True)
        print('\n[4/4] 完成！保存目录：', OUTPUT_DIR)
        print('请核对最新期号是否与网站显示一致；旧期数据不保证网站绝对准确。')
        logging.info('成功导出 %d 期，最新期号 %s', len(records), records[-1]['期号'])
        return 0
    except Exception:
        errors.append(('主程序', traceback.format_exc()))
        print('\n下载失败，原有文件不会因抓取失败而被覆盖。')
        print('最后的详细错误：\n', errors[-1][1][-1800:])
        return 1
    finally:
        for label, detail in errors:
            logging.error('%s\n%s', label, detail)
        print('\n日志：', log_path)


if __name__ == '__main__':
    sys.exit(main())
